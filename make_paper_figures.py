#!/usr/bin/env python
"""Two qualitative figures for the paper, built from runs that already happened.

    python make_paper_figures.py --figure comparison   # Figure A: methods side by side
    python make_paper_figures.py --figure trajectory   # Figure B: RHSO stage by stage

Figure A reads reconstructions that a completed run wrote to disk; it loads no model and
touches no GPU.  Ground truth and the displayed measurement are rebuilt through the same
`ProblemStore` the run used, from the same local pool, so they are the identical arrays --
verified by recomputing PSNR from the saved uint8 and matching the recorded value to
0.005 dB.

Figure B needs states that no run records, so it re-executes RHSO once with
`src.rhso.TRACE` installed (see that module: the collector is write-only, so a traced run
performs the identical sequence of operations as an untraced one).

Both write a vector PDF sized for NeurIPS \\textwidth = 5.5in, plus a PNG preview.
"""
import argparse
import csv
import json
import shutil
import tempfile
from pathlib import Path

from typing import Dict

import numpy as np

CANONICAL_TEXTWIDTH_IN = 5.5          # NeurIPS \textwidth

METHOD_TITLES = {
    "pnp": "PnP-Flow", "dflow": "D-Flow", "mpc_rhc": "MPC-RHC",
    "mpc_delta_t": r"MPC-$\Delta t$", "rhso": "RHSO (ours)", "sdedit": "SDEdit",
}
OURS = "rhso"
OURS_COLOUR = "#137a3a"
PROBLEM_TITLES = {
    "denoising": "Denoising", "deblur": "Deblurring",
    "super_resolution": r"2$\times$ super-res.", "random_inpaint": "Random inpainting",
    "box_inpaint": "Box inpainting", "stroke_painting": "Stroke painting",
}


# =====================================================================================
# Shared: rebuild the problems, and index what a finished run wrote
# =====================================================================================
def build_problems(config_path: str, pool: str):
    """The run's own problem construction, with the data source pinned to a local pool."""
    from src.config import load_config, resolve_run_plan
    from src.data import DataManager
    from src.problems import ProblemStore
    config = load_config(config_path)
    config["data"] = {"source": "local_folder", "local_folder": pool, "image_size": 256}
    plan = resolve_run_plan(config, [], run_id="figures")
    store = ProblemStore()
    store.build_all(plan.problems, DataManager(config, Path("cache") / "data"), verbose=False)
    return {p.name: p for _, p in store.items()}, config, plan


def load_run(run_dir: Path):
    """(model, problem, method) -> the lowest-LPIPS job, with its per-image metrics."""
    rows = list(csv.DictReader(open(run_dir / "results.csv")))
    per = list(csv.DictReader(open(run_dir / "results_per_image.csv")))
    best = {}
    for r in rows:
        key = (r["model"], r["problem"], r["method"])
        try:
            score = float(r["lpips"])
        except (TypeError, ValueError):
            continue
        if key not in best or score < best[key][0]:
            best[key] = (score, r["job_id"])
    metrics = {}
    for r in per:
        metrics.setdefault(r["job_id"], {})[r["image_id"]] = r
    return {k: v[1] for k, v in best.items()}, metrics


_JOB_DIRS: Dict[Path, Dict[str, Path]] = {}


def find_job_dir(run_dir: Path, job_id: str) -> Path:
    """One walk of the run tree, cached: a 100-image run holds 108 job directories."""
    if run_dir not in _JOB_DIRS:
        _JOB_DIRS[run_dir] = {json.load(open(m)).get("job_id"): m.parent
                              for m in run_dir.rglob("metadata.json")}
    out = _JOB_DIRS[run_dir].get(job_id)
    if out is None:
        raise FileNotFoundError("no output directory carries job_id %s" % job_id)
    return out


def reconstruction(job_dir: Path, index: int) -> np.ndarray:
    return np.load(job_dir / "results.npz")["reconstruction"][index]


def to_display(arr: np.ndarray) -> np.ndarray:
    """Canonical [-1,1] float or uint8 -> uint8 RGB for imshow."""
    a = np.asarray(arr)
    if a.dtype == np.uint8:
        return a
    return np.clip((a + 1.0) * 127.5 + 0.5, 0, 255).astype(np.uint8)


def to_canonical(arr: np.ndarray) -> np.ndarray:
    """uint8 back to canonical [-1,1]; the trace cache is stored quantised."""
    a = np.asarray(arr)
    return a.astype(np.float32) / 127.5 - 1.0 if a.dtype == np.uint8 else a.astype(np.float32)


def style():
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    plt.rcParams.update({
        "pdf.fonttype": 42, "ps.fonttype": 42,          # embed real glyphs, not type-3
        "font.family": "serif",
        "font.serif": ["DejaVu Serif", "Times New Roman", "Nimbus Roman"],
        "mathtext.fontset": "dejavuserif",
        "savefig.bbox": "tight", "savefig.pad_inches": 0.01,
        "axes.linewidth": 0.4,
    })
    return plt


def save(fig, out: Path, dpi: int = 400):
    out.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(out, dpi=dpi)
    fig.savefig(out.with_suffix(".png"), dpi=180)
    print("wrote %s  (+ %s)" % (out, out.with_suffix(".png").name))


# =====================================================================================
# Figure A -- one row per inverse problem, one column per method
# =====================================================================================
# Chosen by `--auto-select`: among images where RHSO is best on BOTH LPIPS and PSNR
# against the four rivals at each one's own tuned setting, the largest LPIPS margin, with
# no image reused between rows.  Requiring PSNR too matters -- box inpainting has an image
# whose LPIPS margin is ten times larger but whose PSNR is 1.6 dB WORSE than MPC-RHC's,
# and printing that number under a panel titled RHSO would argue against the figure.
# On the 100-image run.  `--best-available` alone is not good enough here: pure argmax on
# the margin picks images whose statistics are unusual rather than whose reconstruction is
# legible -- a mostly-white geyser, an extreme close-up of fur, a product photo covered in
# text -- where RHSO's lead is large in dB and invisible on the page.  So the shortlist is
# the top both-metric wins per problem and the choice among them is by eye, for a
# recognisable subject with detail in the range each degradation actually damages.
SELECTION = {
    "denoising":        "084_29_peacock",      # +0.108 LPIPS, +1.03 dB; fine feather texture
    "deblur":           "181_48_class181",     # +0.078, +0.81;  curly fur, sharp/soft is obvious
    "super_resolution": "387_68_class387",     # +0.066, +1.06;  red panda, facial detail
    "random_inpaint":   "153_95_class153",     # +0.062, +1.22
    "box_inpaint":      "143_76_class143",     # +0.104, +6.46
}
COMPARISON_METHODS = ["pnp", "dflow", "mpc_rhc", "mpc_delta_t", "rhso"]


def rank_examples(problems, best, metrics, model, methods):
    """(problem -> ranked [(margin, image_id, ...)]) on LPIPS against the best rival."""
    rivals = [m for m in methods if m != "rhso"]
    out = {}
    for problem in SELECTION:
        table = {}
        for method in methods:
            job = best.get((model, problem, method))
            if job:
                table[method] = metrics[job]
        if "rhso" not in table:
            continue
        ranked = []
        for image_id, row in table["rhso"].items():
            present = [m for m in rivals if m in table and image_id in table[m]]
            if not present:
                continue
            lpips = min(float(table[m][image_id]["lpips"]) for m in present)
            psnr = max(float(table[m][image_id]["psnr"]) for m in present)
            both = float(row["lpips"]) < lpips and float(row["psnr"]) > psnr
            ranked.append((both, lpips - float(row["lpips"]),
                           float(row["psnr"]) - psnr, image_id))
        out[problem] = sorted(ranked, key=lambda t: (not t[0], -t[1]))
    return out


def auto_selection(ranking):
    """One image per problem: a both-metric win where possible, never reusing an image."""
    used, chosen = set(), {}
    order = sorted(ranking, key=lambda p: sum(1 for e in ranking[p] if e[0]))
    for problem in order:                      # the most constrained problem picks first
        for both, d_lpips, d_psnr, image_id in ranking[problem]:
            if image_id not in used:
                chosen[problem], _ = image_id, used.add(image_id)
                break
    return chosen


def figure_comparison(args):
    plt = style()
    problems, _, _ = build_problems(args.config, args.pool)
    run_dir = Path(args.run)
    best, metrics = load_run(run_dir)

    ranking = rank_examples(problems, best, metrics, args.model, COMPARISON_METHODS)
    if args.auto_select:
        for problem, ranked in ranking.items():
            print("\n%s" % problem)
            for both, d_lpips, d_psnr, image_id in ranked[:5]:
                print("   %s  dLPIPS %+0.4f  dPSNR %+5.2f  %s"
                      % ("BOTH" if both else "    ", d_lpips, d_psnr, image_id))
        print("\nauto selection: %s" % json.dumps(auto_selection(ranking), indent=2))
        return
    chosen = auto_selection(ranking) if args.best_available else dict(SELECTION)

    def column_title(method):
        """Name D-Flow's trajectory length: `steps` is what separates the one-step
        experiment from a multi-step one, and it is not visible anywhere else."""
        if method != "dflow":
            return METHOD_TITLES[method]
        counts = {int(float(metrics[best[(args.model, p, method)]][chosen[p]]["steps"]))
                  for p in SELECTION}
        if len(counts) != 1:
            return METHOD_TITLES[method]
        n = counts.pop()
        return "%s (%d step%s)" % (METHOD_TITLES[method], n, "" if n == 1 else "s")

    rows = list(SELECTION)
    ncol = 2 + len(COMPARISON_METHODS)
    cell = CANONICAL_TEXTWIDTH_IN / ncol
    fig, axes = plt.subplots(len(rows), ncol,
                             figsize=(CANONICAL_TEXTWIDTH_IN, cell * len(rows) * 1.16))
    fig.subplots_adjust(wspace=0.03, hspace=0.30, left=0.045, right=0.999,
                        top=0.958, bottom=0.028)

    for r, problem in enumerate(rows):
        prob = problems[problem]
        image_id = chosen[problem]
        index = list(prob.image_ids).index(image_id)

        panels = [("Ground truth", to_display(prob.ground_truth[index]), None),
                  ("Measurement", to_display(prob.display_measurement[index]), None)]
        for method in COMPARISON_METHODS:
            job = best[(args.model, problem, method)]
            row = metrics[job][image_id]
            panels.append((column_title(method),
                           to_display(reconstruction(find_job_dir(run_dir, job), index)),
                           float(row["psnr"])))
        degraded = float(metrics[best[(args.model, problem, "rhso")]][image_id]
                         ["degraded_psnr"])
        panels[1] = (panels[1][0], panels[1][1], degraded)

        for c, (title, image, psnr) in enumerate(panels):
            ax = axes[r, c]
            ax.imshow(image, interpolation="lanczos")
            ax.set_xticks([]); ax.set_yticks([])
            winner = title.startswith(METHOD_TITLES[OURS])
            for side in ax.spines.values():
                side.set_linewidth(1.1 if winner else 0.3)
                side.set_color(OURS_COLOUR if winner else "#b8b8b8")
            if r == 0:
                ax.set_title(title, fontsize=6.2, pad=2.6,
                             color=OURS_COLOUR if winner else "black",
                             fontweight="bold" if winner else "normal")
            if c == 0:
                ax.set_ylabel(PROBLEM_TITLES[problem], fontsize=6.0, labelpad=2.2)
            if psnr is not None:
                ax.set_xlabel("PSNR: %.2f" % psnr, fontsize=5.6, labelpad=1.6,
                              color=OURS_COLOUR if winner else "black",
                              fontweight="bold" if winner else "normal")

    save(fig, Path(args.out) / "fig_qualitative_comparison.pdf")


# =====================================================================================
# Figure B -- what RHSO does between pure noise and the reconstruction
# =====================================================================================
def collect_traces(args) -> Path:
    """Re-run RHSO once per traced job with `src.rhso.TRACE` installed, and cache it.

    The states are what makes the figure; they exist only while the strategy runs, so this
    is the one part of the figure pipeline that needs the model and the GPU.  The result is
    an .npz, so redrawing never repeats it.
    """
    import src.rhso as rhso_module
    from src.config import load_config, resolve_run_plan
    from src.data import DataManager
    from src.models import ModelManager, load_adapters
    from src.problems import ProblemStore
    from src.utils import detect_accelerator
    from run import ensure_repositories, init_frameworks, run_job

    config = load_config(args.trajectory_config)
    plan = resolve_run_plan(config, [], run_id="figures")
    cache_root = Path(config["runtime"].get("cache_root", "cache"))
    accel = detect_accelerator(config["runtime"].get("accelerator", "auto"))

    store = ProblemStore()
    store.build_all(plan.problems, DataManager(config, cache_root / "data"), verbose=False)

    context = ensure_repositories(plan, cache_root, verbose=False)
    context.update(init_frameworks(plan, config, accel, cache_root, verbose=False))
    context["ckpt_cache"] = cache_root / "checkpoints"
    load_adapters(plan.resources.models)
    manager = ModelManager(config, plan, False, context)

    saved = {}
    for spec in plan.specs:
        adapter = manager.acquire(spec.model)
        problem = store.get(spec.problem_key)
        print("  tracing %s | %s | N=%d" % (spec.problem, spec.model,
                                            int(spec.num_rhso_steps)))
        rhso_module.TRACE = []
        try:
            run = run_job(adapter, spec, problem, manager)
            trace = rhso_module.TRACE
        finally:
            rhso_module.TRACE = None
        tag = "%s|%s|N=%d" % (spec.problem, spec.model, int(spec.num_rhso_steps))
        # uint8, as every saved reconstruction in this repository already is: the arrays
        # are float32 pixel batches and keeping them so makes the cache gigabytes
        def quantise(a):
            return np.clip((np.asarray(a, np.float32) + 1.0) * 127.5 + 0.5,
                           0, 255).astype(np.uint8)
        saved[tag + "|final"] = quantise(run["pixels"])
        saved[tag + "|s"] = np.asarray([r["s"] for r in trace], np.float32)
        for field in ("state_in", "state_out", "predicted_before", "predicted_after"):
            saved[tag + "|" + field] = np.stack([quantise(r[field]) for r in trace])
    manager.release_all()

    out = Path("cache") / "rhso_traces.npz"
    out.parent.mkdir(parents=True, exist_ok=True)
    np.savez_compressed(out, **saved)
    print("wrote %s  (%d arrays)" % (out, len(saved)))
    return out


def load_traces(args) -> Path:
    # a cache, not a deliverable: it is the model's own intermediate states, regenerable
    # with --collect, and far too large to live beside the figures it feeds
    out = Path("cache") / "rhso_traces.npz"
    if args.collect or not out.exists():
        return collect_traces(args)
    return out


def contact_sheet(args):
    """Every traced array, unstyled -- for choosing the image and the layout."""
    plt = style()
    z = np.load(load_traces(args))
    tags = sorted({k.rsplit("|", 1)[0] for k in z.files})
    for tag in tags:
        n = z[tag + "|state_in"].shape[0]
        images = z[tag + "|state_in"].shape[1]
        fig, axes = plt.subplots(4, n + 1, figsize=(1.1 * (n + 1), 4.6))
        for c in range(n):
            for r, field in enumerate(("state_in", "state_out", "predicted_before",
                                       "predicted_after")):
                ax = axes[r, c]
                ax.imshow(to_display(z[tag + "|" + field][c, args.index]))
                ax.set_xticks([]); ax.set_yticks([])
                if c == 0:
                    ax.set_ylabel(field, fontsize=5)
                if r == 0:
                    ax.set_title("s=%.2f" % z[tag + "|s"][c], fontsize=5)
        for r in range(4):
            axes[r, n].imshow(to_display(z[tag + "|final"][args.index]))
            axes[r, n].set_xticks([]); axes[r, n].set_yticks([])
        axes[0, n].set_title("final", fontsize=5)
        fig.suptitle("%s  image %d of %d" % (tag, args.index, images), fontsize=6)
        save(fig, Path(args.out) / ("contact_%s.pdf" % tag.replace("|", "_")
                                    .replace("=", "")), dpi=110)
        plt.close(fig)


def psnr(a: np.ndarray, b: np.ndarray) -> float:
    return float(10.0 * np.log10(4.0 / np.mean((a - b) ** 2)))


def figure_trajectory(args):
    """One column per RHSO stage: the state, and what the model predicts of it.

    Column k holds everything about stage k and nothing about any other stage --

        row 1   x_{t_k}, the committed state the stage begins from
        row 2   the endpoint the model predicts FROM that state, before optimising it
        row 3   the endpoint it predicts after the stage's inner optimisation

    -- so vertical association carries the whole method: this state, this is where it is
    headed, this is where optimising it sends it instead.  Horizontal arrows across row 1
    are the committed trajectory x_{t_0} -> ... -> x_{t_N}, whose last panel is the
    reconstruction; the vertical arrow inside each column is the one move RHSO makes.

    NOTE ON NOTATION.  The manuscript writes the flow time as t and this repository's code
    writes it as s; they are the same variable and the same numbers.  The figure follows
    the manuscript.
    """
    if args.contact_sheet:
        return contact_sheet(args)
    from matplotlib.patches import FancyArrowPatch
    plt = style()
    z = np.load(load_traces(args))
    problems, _, _ = build_problems(args.trajectory_config, args.pool)

    tag = "%s|%s|N=%d" % (args.problem, args.model, args.n_steps)
    i = args.index
    grid = z[tag + "|s"]                      # canonical time; the manuscript calls it t
    n = len(grid)
    states = [to_canonical(z[tag + "|state_in"][k, i]) for k in range(n)]
    states.append(to_canonical(z[tag + "|state_out"][n - 1, i]))
    before = [to_canonical(z[tag + "|predicted_before"][k, i]) for k in range(n)]
    after = [to_canonical(z[tag + "|predicted_after"][k, i]) for k in range(n)]
    prob = problems[args.problem]
    truth, observed = prob.ground_truth[i], prob.display_measurement[i]
    assert np.allclose(after[-1], states[-1], atol=6e-3)

    # -- geometry -----------------------------------------------------------------------
    # n stage columns + the committed endpoint + a detached reference column
    L, R, TOP, BOT = 0.148, 0.998, 0.878, 0.050
    GAP_UNITS = 0.30
    units = (n + 1) + GAP_UNITS + 1.0
    w = (R - L) / units
    gapx = w * 0.10
    pw = w - gapx
    RG1, RG2 = 0.20, 0.60                     # row 1 -> 2, and row 2 -> 3 (PSNR + arrow)

    height_in = CANONICAL_TEXTWIDTH_IN * pw * (3 + RG1 + RG2) / (TOP - BOT)
    fig = plt.figure(figsize=(CANONICAL_TEXTWIDTH_IN, height_in))
    ph = pw * CANONICAL_TEXTWIDTH_IN / height_in
    rg1, rg2 = RG1 * ph, RG2 * ph
    y_row = [TOP - ph, TOP - 2 * ph - rg1, TOP - 3 * ph - rg1 - rg2]

    def panel(x, y, image, frame="#b8b8b8", lw=0.3):
        ax = fig.add_axes([x, y, pw, ph])
        ax.imshow(to_display(image), interpolation="lanczos")
        ax.set_xticks([]); ax.set_yticks([])
        for side in ax.spines.values():
            side.set_linewidth(lw); side.set_color(frame)
        return ax

    def label(ax, text, colour="black", bold=False, lines=1):
        ax.set_xlabel(text, fontsize=5.6, labelpad=1.7, color=colour, linespacing=1.3,
                      fontweight="bold" if bold else "normal")

    # -- row 1: the committed trajectory ------------------------------------------------
    for c, image in enumerate(states):
        final = (c == n)
        ax = panel(L + c * w, y_row[0], image,
                   OURS_COLOUR if final else "#b8b8b8", 1.1 if final else 0.3)
        ax.set_title("%s\n$x_{t_%d}$\n$t = %.2f$"
                     % ("reconstruction" if final else "stage %d" % (c + 1), c,
                        1.0 if final else grid[c]),
                     fontsize=6.0, pad=2.8, linespacing=1.35,
                     color=OURS_COLOUR if final else "#666",
                     fontweight="bold" if final else "normal")
        if final:
            label(ax, "PSNR: %.2f" % psnr(image, truth), OURS_COLOUR, bold=True)
        if c:
            fig.add_artist(FancyArrowPatch(
                (L + c * w - gapx * 0.92, y_row[0] + ph / 2),
                (L + c * w - gapx * 0.12, y_row[0] + ph / 2), transform=fig.transFigure,
                arrowstyle="-|>", mutation_scale=6, lw=0.8, color="#7a7a7a"))

    # -- rows 2 and 3: the same column, so the association needs no explaining -----------
    for k in range(n):
        x = L + k * w
        label(panel(x, y_row[1], before[k]), "PSNR: %.2f" % psnr(before[k], truth))
        label(panel(x, y_row[2], after[k]),
              "PSNR: %.2f\n(%+.2f)" % (psnr(after[k], truth),
                                       psnr(after[k], truth) - psnr(before[k], truth)))
        fig.add_artist(FancyArrowPatch(
            (x + pw / 2, y_row[2] + ph + rg2 * 0.56), (x + pw / 2, y_row[2] + ph + rg2 * 0.10),
            transform=fig.transFigure, arrowstyle="-|>", mutation_scale=6, lw=0.9,
            color=OURS_COLOUR))

    # -- reference: aligned to rows 1 and 2, so x* sits level with the reconstruction ----
    ref_x = L + (n + 1 + GAP_UNITS) * w
    for j, (image, caption) in enumerate(((truth, "ground truth $x^{\\star}$"),
                                          (observed, "measurement $y$"))):
        ax = panel(ref_x, y_row[j], image)
        label(ax, caption, "#555")
        if j == 0:
            ax.set_title("reference", fontsize=6.0, pad=2.8, color="#555")

    # -- row labels ---------------------------------------------------------------------
    for r, text in enumerate(("state\n$x_{t_k}$",
                              "predicted endpoint\nbefore optimization",
                              "predicted endpoint\nafter optimization")):
        fig.text(L - 0.012, y_row[r] + ph / 2, text, fontsize=5.8, ha="right", va="center",
                 rotation=90, linespacing=1.3, multialignment="center")
    fig.text(L - 0.012, y_row[2] + ph + rg2 * 0.33, "optimize", fontsize=5.8, ha="right",
             va="center", color=OURS_COLOUR, style="italic")

    save(fig, Path(args.out) / ("fig_rhso_trajectory_%s.pdf" % args.problem))


# =====================================================================================
# The same trajectory as an editable slide
# =====================================================================================
def pptx_trajectory(args):
    """Figure B as a .pptx whose every panel, label and arrow is a separate shape.

    Not a picture of the figure: the point of asking for a slide is to move things, so the
    images go in one per panel and the text and arrows are native shapes.  The geometry is
    the PDF's, rescaled to 16:9.
    """
    from PIL import Image
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Emu, Inches, Pt

    z = np.load(load_traces(args))
    problems, _, _ = build_problems(args.trajectory_config, args.pool)
    tag = "%s|%s|N=%d" % (args.problem, args.model, args.n_steps)
    i = args.index
    grid = z[tag + "|s"]
    n = len(grid)
    states = [to_canonical(z[tag + "|state_in"][k, i]) for k in range(n)]
    states.append(to_canonical(z[tag + "|state_out"][n - 1, i]))
    before = [to_canonical(z[tag + "|predicted_before"][k, i]) for k in range(n)]
    after = [to_canonical(z[tag + "|predicted_after"][k, i]) for k in range(n)]
    prob = problems[args.problem]
    truth, observed = prob.ground_truth[i], prob.display_measurement[i]

    # PowerPoint copies each picture into the .pptx, so the files on disk are only a
    # handover and are removed once the deck is written.
    scratch = Path(tempfile.mkdtemp(prefix="rhso_panels_"))

    def png(name, image):
        path = scratch / (name + ".png")
        Image.fromarray(to_display(image)).save(path)
        return str(path)

    GREEN, GREY, INK = RGBColor(0x13, 0x7A, 0x3A), RGBColor(0x88, 0x88, 0x88), \
        RGBColor(0x1A, 0x1A, 0x1A)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide = prs.slides.add_slide(prs.slide_layouts[6])          # blank

    # The left strip is not margin: it carries the rotated row labels and the "optimize"
    # caption.  Sizing it as if it were margin puts both off the slide -- a rotated text box
    # turns about its CENTRE, so a box of width pw ends up occupying its own HEIGHT
    # horizontally, centred where it was placed.
    LABEL_W, TOP_Y = 1.05, 0.66
    LEFT = LABEL_W + 0.15
    GAP_UNITS = 0.30
    w = (13.333 - LEFT - 0.30) / ((n + 1) + GAP_UNITS + 1.0)
    gapx = w * 0.10
    pw = w - gapx
    rg1, rg2 = 0.30, 0.58
    y_row = [TOP_Y, TOP_Y + pw + rg1, TOP_Y + 2 * pw + rg1 + rg2]

    def text(x, y, width, height, lines, size=12, colour=INK, bold=False, italic=False,
             align=PP_ALIGN.CENTER):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = box.text_frame
        frame.word_wrap = False
        frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
        frame.vertical_anchor = MSO_ANCHOR.MIDDLE
        for j, line in enumerate(lines):
            para = frame.paragraphs[0] if j == 0 else frame.add_paragraph()
            para.alignment = align
            run = para.add_run(); run.text = line
            run.font.size = Pt(size); run.font.bold = bold; run.font.italic = italic
            run.font.color.rgb = colour
            run.font.name = "Calibri"
        return box

    def picture(x, y, path, outline=None):
        pic = slide.shapes.add_picture(path, Inches(x), Inches(y), Inches(pw), Inches(pw))
        pic.line.color.rgb = outline or RGBColor(0xB8, 0xB8, 0xB8)
        pic.line.width = Pt(2.0 if outline else 0.75)
        return pic

    def arrow(shape, x, y, width, height, colour):
        sh = slide.shapes.add_shape(shape, Inches(x), Inches(y), Inches(width),
                                    Inches(height))
        sh.fill.solid(); sh.fill.fore_color.rgb = colour
        sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    # -- row 1: the committed trajectory ------------------------------------------------
    for c, image in enumerate(states):
        final = (c == n)
        x = LEFT + c * w
        picture(x, y_row[0], png("state_%d" % c, image), GREEN if final else None)
        text(x, TOP_Y - 0.60, pw, 0.56,
             ["reconstruction" if final else "stage %d" % (c + 1),
              "x(t%d)    t = %.2f" % (c, 1.0 if final else grid[c])],
             size=12, colour=GREEN if final else GREY, bold=final)
        if final:
            text(x, y_row[0] + pw + 0.04, pw, 0.22,
                 ["PSNR: %.2f" % psnr(image, truth)], size=12, colour=GREEN, bold=True)
        if c:
            arrow(MSO_SHAPE.RIGHT_ARROW, x - gapx - 0.03, y_row[0] + pw / 2 - 0.055,
                  gapx + 0.06, 0.11, GREY)

    # -- rows 2 and 3: one stage per column ---------------------------------------------
    for k in range(n):
        x = LEFT + k * w
        picture(x, y_row[1], png("before_%d" % k, before[k]))
        picture(x, y_row[2], png("after_%d" % k, after[k]))
        text(x, y_row[1] + pw + 0.04, pw, 0.22,
             ["PSNR: %.2f" % psnr(before[k], truth)], size=11)
        text(x, y_row[2] + pw + 0.04, pw, 0.40,
             ["PSNR: %.2f" % psnr(after[k], truth),
              "(%+.2f)" % (psnr(after[k], truth) - psnr(before[k], truth))], size=11)
        arrow(MSO_SHAPE.DOWN_ARROW, x + pw / 2 - 0.075, y_row[1] + pw + 0.30,
              0.15, rg2 - 0.34, GREEN)

    # -- reference, level with rows 1 and 2 ---------------------------------------------
    ref_x = LEFT + (n + 1 + GAP_UNITS) * w
    for j, (image, caption) in enumerate(((truth, "ground truth  x*"),
                                          (observed, "measurement  y"))):
        picture(ref_x, y_row[j], png("ref_%d" % j, image))
        text(ref_x, y_row[j] + pw + 0.04, pw, 0.22, [caption], size=11, colour=GREY)
        if j == 0:
            text(ref_x, TOP_Y - 0.32, pw, 0.24, ["reference"], size=12, colour=GREY)

    # -- row labels, rotated up the left edge -------------------------------------------
    for r, lines in enumerate((["state  x(tk)"],
                               ["predicted endpoint", "before optimization"],
                               ["predicted endpoint", "after optimization"])):
        box = text(LEFT - 0.30 - pw / 2, y_row[r] + pw / 2 - 0.20, pw, 0.40, lines, size=11)
        box.rotation = 270
    text(LEFT - 1.05, y_row[1] + pw + 0.30, 0.95, rg2 - 0.34, ["optimize"], size=11,
         colour=GREEN, italic=True, align=PP_ALIGN.RIGHT)

    # No renderer here to look at the result, so check the geometry instead: a shape off
    # the slide is invisible in PowerPoint and silent in python-pptx.
    W, H = 13.333, 7.5
    for sh in slide.shapes:
        x0, y0 = sh.left / 914400.0, sh.top / 914400.0
        x1, y1 = x0 + sh.width / 914400.0, y0 + sh.height / 914400.0
        if sh.rotation:                       # rotation is about the centre
            cx, cy = (x0 + x1) / 2, (y0 + y1) / 2
            half_w, half_h = (y1 - y0) / 2, (x1 - x0) / 2
            x0, x1, y0, y1 = cx - half_w, cx + half_w, cy - half_h, cy + half_h
        if x0 < -0.01 or y0 < -0.01 or x1 > W + 0.01 or y1 > H + 0.01:
            raise SystemExit("shape %r lands outside the slide: x %.2f..%.2f  y %.2f..%.2f"
                             % (sh.shape_type, x0, x1, y0, y1))

    out = Path(args.out) / ("fig_rhso_trajectory_%s.pptx" % args.problem)
    prs.save(out)
    shutil.rmtree(scratch, ignore_errors=True)
    print("wrote %s  (%d shapes, %.2f x %.2f in panels)" % (out, len(slide.shapes), pw, pw))


def main():
    p = argparse.ArgumentParser(description=__doc__,
                                formatter_class=argparse.RawDescriptionHelpFormatter)
    p.add_argument("--figure", choices=("comparison", "trajectory"), required=True)
    p.add_argument("--run", default="outputs/t1_hpo", help="a finished run directory")
    p.add_argument("--config", default="configs/experiments_t1_hpo.yaml")
    p.add_argument("--pool", default="cache/data/imagenet_val_100")
    p.add_argument("--model", default="pmf")
    p.add_argument("--out", default="figures")
    p.add_argument("--auto-select", action="store_true",
                   help="print the example ranking instead of drawing")
    p.add_argument("--best-available", action="store_true",
                   help="draw the top-ranked image per problem instead of SELECTION")
    p.add_argument("--trajectory-config", default="configs/figure_trajectory.yaml")
    p.add_argument("--pptx", action="store_true",
                   help="write the trajectory as an editable slide instead of a PDF")
    p.add_argument("--contact-sheet", action="store_true",
                   help="dump every traced array so a layout can be chosen")
    p.add_argument("--index", type=int, default=6, help="which image of the traced batch")
    p.add_argument("--problem", default="random_inpaint")
    p.add_argument("--n-steps", type=int, default=4, help="the traced num_rhso_steps")
    p.add_argument("--collect", action="store_true",
                   help="re-run the traced RHSO jobs (needs the GPU) before drawing")
    args = p.parse_args()
    if args.figure == "comparison":
        figure_comparison(args)
    elif args.pptx:
        pptx_trajectory(args)
    else:
        figure_trajectory(args)


if __name__ == "__main__":
    main()
